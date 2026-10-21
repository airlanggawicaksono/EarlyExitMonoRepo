"""Smoke test of exp2/analyze/update_pack.py on a tiny synthetic input (no model, no board)."""
import csv
import json

from openpyxl import load_workbook

from exp2.analyze import update_pack as U


def _w(path, cols, rows):
    with open(path, "w", newline="") as fh:
        wr = csv.writer(fh)
        wr.writerow(cols)
        wr.writerows(rows)


def _session(d, fam, state, spawn, ready):
    s = d / f"session_run1c_{fam}_{state}"
    (s / "telemetry").mkdir(parents=True)
    sess = {"servers": {fam: {"t_spawn_ns": spawn, "t_ready_ns": ready,
                              "ready": {"warmup": {"load_s": 1.0, "first_infer_s": {"5": 2.0, "23": 0.5}}}}}}
    (s / "session.json").write_text(json.dumps(sess))
    _w(s / "telemetry" / "power.csv", ["t_mono_ns", "vdd_in_mw"],
       [[t * 10**9, 10000] for t in range(0, 12)])
    return s


def test_pack_builds_and_reloads(tmp_path):
    d = tmp_path / "csv"
    d.mkdir()
    _w(d / "solo.csv", ["type_id", "f_mhz", "repeat", "n", "tau_p50", "tau_p95", "p_board_w", "window_s"],
       [[t, f, r, 100, 0.1 * (1 + r * 0.01), 0.12, 10.0, 10.0] for t in ("bert:5", "llama:3")
        for f in (612, 816) for r in (0, 1)])
    cols = ["cell_id", "split", "f_mhz", "repeat", "type_id", "family", "A", "n", "n_window", "s_p50", "s_p95",
            "tau_solo_p50", "window_s", "energy_j", "p_board_w"]
    rows = []
    for f in (612, 816):
        for r in (0, 1):
            for t, fam in (("bert:5", "bert"), ("llama:3", "llama")):
                rows.append(["c001", "train", f, r, t, fam, "bert:5|llama:3", 50, 40, 2.0, 2.2, 0.1, 20.0, 300.0,
                             15.0])
    _w(d / "corun.csv", cols, rows)
    _w(d / "idle.csv", ["f_mhz", "kind", "p_idle_w"], [[612, "resident", 5.0], [816, "resident", 5.0]])
    _w(d / "warmup.csv", ["session_id", "family", "t_load_s", "first_infer_s_max", "warm_energy_j", "cache"],
       [["s1", "bert", 16.0, 0.7, 44.0, "warm"]])
    s1 = [_session(tmp_path, "bert", "cold", 0, 8 * 10**9), _session(tmp_path, "bert", "warm", 0, 4 * 10**9)]
    figs = tmp_path / "figs"
    figs.mkdir()
    out = tmp_path / "update_17"
    lines, counts = U.build(d, s1, figs, out)
    assert (out / "README.md").is_file()
    assert (out / "1_cost_value" / "fig_cost_delta.png").is_file()
    assert (out / "3_slowdown_factor" / "fig_slowdown_by_n.png").is_file()
    wb = load_workbook(out / "1_cost_value" / "cost_bert.xlsx")
    assert wb.sheetnames == ["summary", "bert_5", "bert_23", "definitions"]
    ws = wb["bert_5"]
    hdr = [c.value for c in ws[1]]
    med = [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)
           if r[hdr.index("repeat")] == "median" and r[hdr.index("n_concurrent")] == 2]
    assert len(med) == 2
    m = [x for x in med if x["f_mhz"] == 612][0]
    assert abs(m["tau_p50_s"] - 0.2) < 1e-9 and abs(m["delta_tau_pct"] - 100.0) < 1e-9
    assert abs(m["energy_total_per_request_j"] - 300.0 / 8 * 0.1) < 1e-9    # e_work = 300 / (2 x 40 x 0.1 s), x tau_solo 0.1
    assert abs(m["delta_energy_total_pct"] - 275.0) < 1e-6                  # solo: 10 W x 10 s / 100 = 1 J
    assert abs(m["energy_dyn_per_request_j"] - (300.0 - 5.0 * 20.0) / 8 * 0.1) < 1e-9   # idle 5 W, window 20 s
    assert abs(m["delta_energy_dyn_pct"] - 400.0) < 1e-6                    # solo dynamic: 5 W x 10 s / 100 = 0.5 J
    sl = load_workbook(out / "3_slowdown_factor" / "slowdown_llama.xlsx")["llama_3"]
    h = [c.value for c in sl[1]]
    cap = {r[h.index("capacity_c")] for r in sl.iter_rows(min_row=2, values_only=True)}
    assert cap == {1.0}                                           # 1/2 + 1/2
    wsu = load_workbook(out / "2_warmup" / "warmup_bert.xlsx")["summary"]
    first = [c.value for c in wsu[2]]
    assert abs(first[4] - (-4.0)) < 1e-9                          # warm 4 s minus cold 8 s
    assert all(isinstance(n, int) for c in counts.values() for n in c.values())
