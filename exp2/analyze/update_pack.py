"""Results pack "update_17": figures and one Excel workbook per model family, solo versus n concurrent.

Three folders: 1_cost_value, 2_warmup, 3_slowdown_factor. Workbooks: one per family (bert, llama, yolo), one sheet per
request type plus "summary" and "definitions". matplotlib, numpy, openpyxl and the standard library only. No model run.
Reuses the helpers of exp2/analyze/figures.py (style, save, loaders, capacity_table, energy integral).
CLI: python -m exp2.analyze.update_pack --run1-csv DIR --run1c DIR [DIR ...] --figures DIR --out DIR
  --run1-csv  exp2/frozen/run2/run1_csv (solo.csv corun.csv warmup.csv)
  --run1c     Run 1c session folders session_run1c_{family}_{cold,warm}
  --figures   folder with the paper figures (PNG + PDF), copied into the pack
  --out       update_17 folder
"""
import argparse
import collections
import shutil
import statistics as st
import sys
from pathlib import Path

import numpy as np
from matplotlib.lines import Line2D
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from exp2.analyze import figures as F

plt = F.plt
FAMS = ("bert", "llama", "yolo")
TYPES_OF = {fam: [t for t in F.TYPES if t.startswith(fam + ":")] for fam in FAMS}
DATE = "2026-10-21"

COST_FIGS = ("fig_solo_time", "fig_power", "fig_gamma", "fig_run2_energy_attain")
WARM_FIGS = ("fig_warmup", "fig3d_warmup")
SLOW_FIGS = ("fig_capacity", "fig3d_capacity", "fig_slowdown_fit")

# paper values (results.tex): tab:startup and tab:capacity / section prose
PAPER_STARTUP = {"bert": (49.6, 21.4, 2.3, 0.6), "llama": (91.1, 72.2, 12.1, 4.1), "yolo": (156.9, 41.3, 72.3, 14.7)}
PAPER_CAP_GROUPS = {   # group -> {f: (min, max)}
    "encoder and detector": {612: (0.90, 0.96), 816: (0.94, 1.01), 1020: (0.97, 1.03)},
    "with llama:3": {612: (1.03, 1.10), 816: (1.10, 1.18), 1020: (1.13, 1.20)},
    "with llama:15": {612: (0.95, 1.00), 816: (1.09, 1.20), 1020: (1.18, 1.34)},
}


# ------------------------------------------------------------------------------------------------ helpers
def _f(x):
    return float(x)


def _med(v):
    v = [x for x in v if x is not None]
    return st.median(v) if v else None


def _pct(new, ref):
    if new is None or ref in (None, 0):
        return None
    return (new / ref - 1.0) * 100.0


def _diff(new, ref):
    return None if new is None or ref is None else new - ref


def _n_of(A):
    return A.count("|") + 1


def _src(cell_id):
    return "d" if cell_id.startswith("d_") else "c"


def load_inputs(run1_csv):
    d = Path(run1_csv)
    solo = F._csv(d / "solo.csv")
    corun = F._csv(d / "corun.csv")
    try:
        warm = F._csv(d / "warmup.csv")
    except F.Skip:
        warm = []
    try:
        idle = {int(_f(r["f_mhz"])): _f(r["p_idle_w"]) for r in F._csv(d / "idle.csv") if r["kind"] == "resident"}
    except F.Skip:
        idle = {}
    try:
        mem = {(r["family"], r["exit"]): (_f(r["resident_mb"]), _f(r["act_mb"])) for r in F._csv(d / "mem.csv")}
    except F.Skip:
        mem = {}
    try:
        screen = F._csv(d / "screen.csv")
    except F.Skip:
        screen = []
    return solo, corun, warm, idle, mem, screen


def corun_groups(corun):
    """(cell_id, f, repeat) -> member rows of the set (one per tenant)."""
    g = collections.defaultdict(list)
    for r in corun:
        g[(r["cell_id"], int(_f(r["f_mhz"])), r["repeat"])].append(r)
    return g


def _prefer_d(rows_by_src):
    """Run 1d rows replace Run 1c rows where Run 1d measured the same set and level (as capacity_table)."""
    return rows_by_src.get("d") or rows_by_src.get("c") or []


# ------------------------------------------------------------------------------------------------ Excel writer
def write_workbook(path, sheets, definitions):
    """sheets: [(name, columns, rows)], rows are dicts; definitions: [(term, text)]. Header bold and frozen."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(list(cols))
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(12, min(34, len(str(c)) + 3))
            fmt = "0.0000" if (c.endswith("_s") or c.endswith("_w") or c.endswith("_j") or "pct" in c or c in (
                "slowdown_s_p50", "capacity_c", "ratio_warm_cold") or c.startswith("delta")) else None
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
                for x in cell:
                    if fmt and isinstance(x.value, float):
                        x.number_format = fmt
    ws = wb.create_sheet("definitions")
    ws.append(["term", "definition"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for t, d in definitions:
        ws.append([t, d])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------------------------------------ 1 cost value
def set_level_dyn_delta(solo, corun, p_idle):
    """(A, f) -> median over repeats (Run 1d rows where they exist) of the set-level dynamic delta (%):
    (energy_j - P_idle x window_s) / sum_i n_window_i x (solo dynamic energy per request of i) - 1."""
    _, _, ds = _solo_refs(solo, p_idle)
    acc = collections.defaultdict(lambda: collections.defaultdict(list))
    for (cid, f, rep), mem in corun_groups(corun).items():
        num = _f(mem[0]["energy_j"]) - p_idle[f] * _f(mem[0]["window_s"])
        den = sum(_f(x["n_window"]) * ds[(x["type_id"], f)] for x in mem)
        acc[(mem[0]["A"], f)][_src(cid)].append((num / den - 1.0) * 100.0)
    return {k: st.median(_prefer_d(v)) for k, v in acc.items()}


COST_COLS = ["type_id", "n_concurrent", "f_mhz", "co_run_set_A", "cell_id", "split", "repeat", "source",
             "tau_p50_s", "tau_p95_s", "p_board_w", "window_s", "completed_in_window", "energy_window_j",
             "p_idle_w", "e_work_total_j_per_solo_s", "dyn_ratio_D_over_S",
             "energy_total_per_request_j", "energy_dyn_per_request_j", "ref_tau_p50_s", "ref_p_board_w",
             "ref_energy_total_per_request_j", "ref_energy_dyn_per_request_j",
             "delta_tau_s", "delta_tau_pct", "delta_p_w", "delta_p_pct",
             "delta_energy_total_j", "delta_energy_total_pct", "delta_energy_dyn_j", "delta_energy_dyn_pct",
             "resident_mb", "act_mb", "act_sum_mb", "ref_act_mb", "delta_act_mb", "delta_act_pct",
             "min_mem_available_mb", "ref_min_mem_available_mb",
             "delta_mem_available_mb", "delta_mem_available_pct",
             "completed_type_in_window", "throughput_rps", "set_throughput_rps", "ref_throughput_rps",
             "delta_throughput_rps", "delta_throughput_pct",
             "gpu_util_measured", "gpu_demand_offered"]
NUMERIC = ("min_mem_available_mb", "completed_type_in_window",
           "throughput_rps", "set_throughput_rps", "gpu_util_measured", "gpu_demand_offered", "tau_p50_s", "tau_p95_s", "p_board_w", "window_s", "completed_in_window", "energy_window_j", "p_idle_w",
           "e_work_total_j_per_solo_s", "dyn_ratio_D_over_S", "energy_total_per_request_j",
           "energy_dyn_per_request_j")


def _solo_refs(solo, p_idle):
    """(type, f) -> median over repeats of power, total and dynamic energy per request of solo.csv."""
    p, e, d = collections.defaultdict(list), collections.defaultdict(list), collections.defaultdict(list)
    for r in solo:
        f = int(_f(r["f_mhz"]))
        k = (r["type_id"], f)
        p[k].append(_f(r["p_board_w"]))
        e[k].append(_f(r["p_board_w"]) * _f(r["window_s"]) / _f(r["n"]))
        if f in p_idle:
            d[k].append((_f(r["p_board_w"]) - p_idle[f]) * _f(r["window_s"]) / _f(r["n"]))
    return ({k: st.median(v) for k, v in p.items()}, {k: st.median(v) for k, v in e.items()},
            {k: st.median(v) for k, v in d.items()})


def _cost_row(base, refs):
    """refs: dict(tau, p, e, d) = solo reference at the same frequency."""
    r = dict(base)
    r.update(ref_tau_p50_s=refs["tau"], ref_p_board_w=refs["p"], ref_energy_total_per_request_j=refs["e"],
             ref_energy_dyn_per_request_j=refs["d"],
             delta_tau_s=_diff(r["tau_p50_s"], refs["tau"]), delta_tau_pct=_pct(r["tau_p50_s"], refs["tau"]),
             delta_p_w=_diff(r["p_board_w"], refs["p"]), delta_p_pct=_pct(r["p_board_w"], refs["p"]),
             delta_energy_total_j=_diff(r["energy_total_per_request_j"], refs["e"]),
             delta_energy_total_pct=_pct(r["energy_total_per_request_j"], refs["e"]),
             delta_energy_dyn_j=_diff(r["energy_dyn_per_request_j"], refs["d"]),
             delta_energy_dyn_pct=_pct(r["energy_dyn_per_request_j"], refs["d"]))
    r.update(ref_act_mb=refs.get("act"), ref_min_mem_available_mb=refs.get("mem"),
             ref_throughput_rps=refs.get("thr"),
             delta_act_mb=_diff(r.get("act_sum_mb"), refs.get("act")),
             delta_act_pct=_pct(r.get("act_sum_mb"), refs.get("act")),
             delta_mem_available_mb=_diff(r.get("min_mem_available_mb"), refs.get("mem")),
             delta_mem_available_pct=_pct(r.get("min_mem_available_mb"), refs.get("mem")),
             delta_throughput_rps=_diff(r.get("throughput_rps"), refs.get("thr")),
             delta_throughput_pct=_pct(r.get("throughput_rps"), refs.get("thr")))
    return r


STATS = {"matched": 0, "unmatched": 0}


def _num(x):
    return None if x in (None, "") else float(x)


def screen_index(screen):
    """(key, f) -> kept screen rows sorted by (session_id, repeat in the session). key = cell_id (co-run) or the
    type id (solo; 'd_solo_...' conditions of Run 1d join the solo type). Idle conditions are skipped.
    Condition id formats: solo_bert-5_f612_r0, c001_f612_r2, d_c002_f816_r0, d_solo_llama-15_f612_r0."""
    idx = collections.defaultdict(list)
    for x in screen:
        if x["mode"] == "idle" or x["kept"] != "1":
            continue
        base, f, r = x["condition_id"].rsplit("_", 2)
        if base.startswith("d_solo_"):
            base = base[2:]
        if base.startswith("solo_"):
            fam, ex = base[5:].rsplit("-", 1)
            base = f"{fam}:{ex}"
        idx[(base, int(f[1:]))].append((x["session_id"], int(r[1:]), x))
    return {k: [v[2] for v in sorted(v, key=lambda a: (a[0], a[1]))] for k, v in idx.items()}


def _screen_for(sidx, ranks, key, f, repeat):
    """Screen row of (key, f, repeat): the repeat counts of corun.csv / solo.csv are shifted over the sessions, so
    the k-th smallest repeat of (key, f) is the k-th kept screen row in session order (counts agree 1:1)."""
    rows = sidx.get((key, f), [])
    k = ranks.get((key, f), {}).get(repeat)
    if k is None or k >= len(rows):
        STATS["unmatched"] += 1
        return None
    STATS["matched"] += 1
    return rows[k]


def build_cost(solo, corun, p_idle=None, mem=None, screen=None):
    """-> {type: [rows]}: n = 1 rows (solo.csv), n = 2, 3 rows (corun.csv), median-over-repeats rows.
    p_idle: {f: W}, warm idle power with the models resident (idle.csv kind resident = power.json p_idle_w)."""
    p_idle = p_idle or {}
    mem = mem or {}
    for k in STATS:
        STATS[k] = 0
    sidx = screen_index(screen or [])
    ranks = collections.defaultdict(dict)
    for key, f, rp in sorted({(r["type_id"], int(_f(r["f_mhz"])), int(r["repeat"])) for r in solo}):
        ranks[(key, f)][rp] = len(ranks[(key, f)])
    for key, f, rp in sorted({(r["cell_id"], int(_f(r["f_mhz"])), int(r["repeat"])) for r in corun}):
        ranks[(key, f)][rp] = len(ranks[(key, f)])

    def act_of(t):
        fam, ex = t.split(":")
        return mem.get((fam, ex), (None, None))

    def memcols(srow):
        if srow is None:
            return dict(min_mem_available_mb=None)
        return dict(min_mem_available_mb=_num(srow["min_mem_available_mb"]))

    ps, es, ds = _solo_refs(solo, p_idle)
    sm, st_thr, st_u = collections.defaultdict(list), collections.defaultdict(list), collections.defaultdict(list)
    for r in solo:
        k = (r["type_id"], int(_f(r["f_mhz"])))
        st_thr[k].append(_f(r["n"]) / _f(r["window_s"]))
        st_u[k].append(_f(r["u"]))
        v = memcols(_screen_for(sidx, ranks, k[0], k[1], int(r["repeat"])))["min_mem_available_mb"]
        if v is not None:
            sm[k].append(v)
    thr_ref = {k: st.median(v) for k, v in st_thr.items()}
    u_ref = {k: st.median(v) for k, v in st_u.items()}
    mem_ref = {k: st.median(v) for k, v in sm.items()}
    STATS["matched"] = STATS["unmatched"] = 0

    def refs_of(t, f, tau):
        k = (t, f)
        return dict(tau=tau, p=ps.get(k), e=es.get(k), d=ds.get(k), act=act_of(t)[1], mem=mem_ref.get(k),
                    thr=thr_ref.get(k))

    tau_ref = {}
    for r in corun:
        tau_ref[(r["type_id"], int(_f(r["f_mhz"])))] = _f(r["tau_solo_p50"])
    out = collections.defaultdict(list)
    for r in solo:
        t, f = r["type_id"], int(_f(r["f_mhz"]))
        pi = p_idle.get(f)
        n, w, p = _f(r["n"]), _f(r["window_s"]), _f(r["p_board_w"])
        base = dict(type_id=t, n_concurrent=1, f_mhz=f, co_run_set_A=t, cell_id="solo", split="", repeat=int(r["repeat"]),
                    source="solo", tau_p50_s=_f(r["tau_p50"]), tau_p95_s=_f(r["tau_p95"]), p_board_w=p, window_s=w,
                    completed_in_window=int(n), energy_window_j=p * w, p_idle_w=pi, e_work_total_j_per_solo_s=None,
                    dyn_ratio_D_over_S=None, energy_total_per_request_j=p * w / n,
                    energy_dyn_per_request_j=None if pi is None else (p - pi) * w / n)
        base.update(memcols(_screen_for(sidx, ranks, t, f, int(r["repeat"]))))
        base.update(resident_mb=act_of(t)[0], act_mb=act_of(t)[1], act_sum_mb=act_of(t)[1],
                    completed_type_in_window=int(n), throughput_rps=n / w, set_throughput_rps=n / w,
                    gpu_util_measured=_f(r["u"]), gpu_demand_offered=None)
        out[t].append(_cost_row(base, refs_of(t, f, tau_ref.get((t, f)))))
    for (cid, f, rep), members in sorted(corun_groups(corun).items()):
        done = sum(_f(x["n_window"]) for x in members)
        work = sum(_f(x["n_window"]) * _f(x["tau_solo_p50"]) for x in members)     # W_window (s)
        e_tot = _f(members[0]["energy_j"]) / work
        pi = p_idle.get(f)
        dyn = None      # D_window / S_window: dynamic energy of the window / serial-equivalent dynamic energy
        if pi is not None:
            s_win = sum(_f(x["n_window"]) * ds[(x["type_id"], f)] for x in members)
            dyn = (_f(members[0]["energy_j"]) - pi * _f(members[0]["window_s"])) / s_win
        for x in members:
            t = x["type_id"]
            ts = _f(x["tau_solo_p50"])
            base = dict(type_id=t, n_concurrent=_n_of(x["A"]), f_mhz=f, co_run_set_A=x["A"], cell_id=cid,
                        split=x["split"], repeat=int(rep), source=_src(cid), tau_p50_s=_f(x["s_p50"]) * ts,
                        tau_p95_s=_f(x["s_p95"]) * ts, p_board_w=_f(x["p_board_w"]), window_s=_f(x["window_s"]),
                        completed_in_window=int(done), energy_window_j=_f(x["energy_j"]), p_idle_w=pi,
                        e_work_total_j_per_solo_s=e_tot, dyn_ratio_D_over_S=dyn,
                        energy_total_per_request_j=e_tot * ts,
                        energy_dyn_per_request_j=None if dyn is None else ds[(t, f)] * dyn)
            base.update(memcols(_screen_for(sidx, ranks, cid, f, int(rep))))
            acts = [act_of(m["type_id"])[1] for m in members]
            base.update(resident_mb=act_of(t)[0], act_mb=act_of(t)[1],
                        act_sum_mb=None if None in acts else sum(acts),
                        completed_type_in_window=int(_f(x["n_window"])),
                        throughput_rps=_f(x["n_window"]) / _f(x["window_s"]),
                        set_throughput_rps=done / _f(x["window_s"]), gpu_util_measured=None,
                        gpu_demand_offered=_f(x["U_A"]))
            out[t].append(_cost_row(base, refs_of(t, f, ts)))
    for t in list(out):
        out[t].extend(_median_rows(out[t], NUMERIC, _cost_row))
    return out


def _median_rows(rows, numeric, mk):
    """Median over repeats of each (type, set, f): Run 1d rows where they exist (as capacity_table)."""
    grp = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in rows:
        grp[(r["type_id"], r["n_concurrent"], r["co_run_set_A"], r["f_mhz"])][r["source"] if r["source"] in "cd"
                                                                              else "c"].append(r)
    res = []
    for (t, n, A, f), bysrc in sorted(grp.items()):
        rs = _prefer_d(bysrc)
        base = {k: _med([x[k] for x in rs]) for k in numeric}
        base["completed_in_window"] = int(round(base["completed_in_window"]))
        base["completed_type_in_window"] = int(round(base["completed_type_in_window"]))
        for k in ("resident_mb", "act_mb", "act_sum_mb"):
            base[k] = rs[0][k]
        base.update(type_id=t, n_concurrent=n, f_mhz=f, co_run_set_A=A, cell_id="median", split=rs[0]["split"],
                    repeat="median", source=("d" if bysrc.get("d") else "c" if n > 1 else "solo"))
        ref = rs[0]
        res.append(mk(base, dict(tau=ref["ref_tau_p50_s"], p=ref["ref_p_board_w"],
                                 e=ref["ref_energy_total_per_request_j"], d=ref["ref_energy_dyn_per_request_j"],
                                 act=ref["ref_act_mb"], mem=ref["ref_min_mem_available_mb"],
                                 thr=ref["ref_throughput_rps"])))
    return res


def cost_summary(rows_by_type, fam):
    out = []
    for n in (1, 2, 3):
        for f in (408, 612, 714, 816, 1020):
            rs = [r for t in TYPES_OF[fam] for r in rows_by_type.get(t, []) if r["repeat"] == "median"
                  and r["n_concurrent"] == n and r["f_mhz"] == f]
            if not rs:
                continue
            out.append(dict(family=fam, n_concurrent=n, f_mhz=f, median_rows=len(rs),
                            tau_p50_s=_med([r["tau_p50_s"] for r in rs]),
                            p_board_w=_med([r["p_board_w"] for r in rs]),
                            energy_total_per_request_j=_med([r["energy_total_per_request_j"] for r in rs]),
                            energy_dyn_per_request_j=_med([r["energy_dyn_per_request_j"] for r in rs]),
                            delta_tau_pct=_med([r["delta_tau_pct"] for r in rs]),
                            delta_p_pct=_med([r["delta_p_pct"] for r in rs]),
                            delta_energy_total_pct=_med([r["delta_energy_total_pct"] for r in rs]),
                            delta_energy_dyn_pct=_med([r["delta_energy_dyn_pct"] for r in rs]),
                            act_sum_mb=_med([r["act_sum_mb"] for r in rs]),
                            delta_act_mb=_med([r["delta_act_mb"] for r in rs]),
                            min_mem_available_mb=_med([r["min_mem_available_mb"] for r in rs]),
                            delta_mem_available_mb=_med([r["delta_mem_available_mb"] for r in rs]),
                            delta_mem_available_pct=_med([r["delta_mem_available_pct"] for r in rs]),
                            throughput_rps=_med([r["throughput_rps"] for r in rs]),
                            set_throughput_rps=_med([r["set_throughput_rps"] for r in rs]),
                            delta_throughput_rps=_med([r["delta_throughput_rps"] for r in rs]),
                            delta_throughput_pct=_med([r["delta_throughput_pct"] for r in rs]),
                            gpu_util_measured=_med([r["gpu_util_measured"] for r in rs]),
                            gpu_demand_offered=_med([r["gpu_demand_offered"] for r in rs])))
    return out


SUM_COLS = ["family", "n_concurrent", "f_mhz", "median_rows", "tau_p50_s", "p_board_w",
            "energy_total_per_request_j", "energy_dyn_per_request_j", "delta_tau_pct", "delta_p_pct",
            "delta_energy_total_pct", "delta_energy_dyn_pct", "act_sum_mb", "delta_act_mb", "min_mem_available_mb",
            "delta_mem_available_mb", "delta_mem_available_pct", "throughput_rps", "set_throughput_rps",
            "delta_throughput_rps", "delta_throughput_pct", "gpu_util_measured", "gpu_demand_offered"]

COST_DEFS = [
    ("n_concurrent", "Number of requests in service in total: 1 = solo (solo.csv), 2 = pair, 3 = triple (corun.csv, n = number of members of set A)."),
    ("tau_p50_s / tau_p95_s", "Median / 95th percentile service time of the type (s). Solo: solo.csv. Co-run: s_p50 (s_p95) x tau_solo_p50 of corun.csv."),
    ("p_board_w", "Mean board power VDD_IN over the measurement window (W). For a co-run row, the power of the whole set, not of one tenant."),
    ("energy_window_j", "Energy of the window (J). Solo: p_board_w x window_s. Co-run: energy_j of corun.csv (whole set)."),
    ("completed_in_window", "Requests completed in the window. Solo: n of solo.csv. Co-run: sum of n_window over all tenants of the set."),
    ("p_idle_w", "Warm idle board power at the same frequency with the models resident (W): idle.csv kind resident, equal to p_idle_w of power.json (the value gamma uses)."),
    ("e_work_total_j_per_solo_s", "Co-run only: energy_window_j / W_window, W_window = sum over the tenants i of n_window_i x tau_solo_p50_i (solo-equivalent work of the window, s, same frequency)."),
    ("dyn_ratio_D_over_S", "Co-run only: D_window / S_window. D_window = energy_window_j - p_idle_w x window_s. S_window = sum over the tenants i of n_window_i x (solo dynamic energy per request of i at the same frequency). Same for every member of the window; delta_energy_dyn_pct = (ratio - 1) x 100. Comparable to 1/gamma - 1 of the paper."),
    ("energy_total_per_request_j", "Board energy INCLUDING idle, busy period, per request of THIS type. Solo: p_board_w x window_s / n (solo.csv). Co-run: e_work_total x tau_solo_p50 of the type (work-proportional attribution). The mixed-type figure (energy_j / requests of all tenants) is not in this sheet: not comparable to solo."),
    ("energy_dyn_per_request_j", "Dynamic energy per request of this type, idle subtracted. Solo: (p_board_w - p_idle_w) x window_s / n. Co-run: solo dynamic energy per request of the type x dyn_ratio_D_over_S (attribution in proportion to solo dynamic energy)."),
    ("attribution", "The energy of one request type inside a shared GPU is not separately measurable, so every per-type energy value is an attribution. Total uses tau-proportional attribution, dynamic uses solo-dynamic-energy attribution, and dynamic is the one comparable to gamma in the paper."),
    ("why total can be negative", "At 1020 MHz a set has capacity above 1, so it completes more work per second and shares the idle power over more work, and the total energy per request can fall below solo. The dynamic energy has no idle part and rises (gamma below 1 means co-running costs more dynamic energy per unit of work than serial execution), so delta_energy_dyn_pct stays positive."),
    ("ref_*", "Solo reference at the same GPU frequency. tau: tau_solo_p50 of corun.csv (pooled solo median). Power and energy: median over solo repeats of solo.csv."),
    ("delta_X_s / _w / _j", "Concurrent value minus solo reference (same unit as the value). delta_energy_total_j and delta_energy_dyn_j use the two energies per request above."),
    ("delta_X_pct", "(concurrent / solo - 1) x 100. delta_energy_total_pct and delta_energy_dyn_pct likewise."),
    ("repeat = median", "Median over repeats of one (type, set, frequency); Run 1d rows replace Run 1c rows where Run 1d re-measured the cell (source d). Deltas are recomputed from the medians."),
    ("split", "train or heldout cell of the slowdown fit (corun.csv)."),
    ("resident_mb / act_mb", "mem.csv: resident_mb = PyTorch allocator reserved memory of the type's family at ready (MB). act_mb = activation memory of one request of the type (MB)."),
    ("act_sum_mb", "Sum of act_mb over the requests in service: n = 1 the type itself, n = 2 or 3 all members of set A. delta_act_mb / delta_act_pct against the act_mb of the type alone."),
    ("min_mem_available_mb", "Lowest MemAvailable (MB) during the condition (screen.csv). screen.csv has no swap or page-fault values for Run 1 (max_swap_used_mb and pgmajfault_delta are empty in every row), so these columns are not in the workbooks."),
    ("screen.csv join", "condition_id formats: solo_bert-5_f612_r0, c001_f612_r2, d_c002_f816_r0, d_solo_llama-15_f612_r0 (d_ = Run 1d). Key = (cell or type, f). The repeat of solo.csv / corun.csv is shifted over sessions, so the k-th smallest repeat of a key takes the k-th kept screen row in session_id order (the counts agree one to one). The repeat order within a key follows session_id order (an assumption); the medians over repeats in the workbooks do not depend on it. Only kept = 1 rows."),
    ("delta_mem_available_*", "min_mem_available_mb minus the median over solo repeats of the same type and level. Negative = less memory left than solo."),
    ("memory caution", "In Run 1 all three families stay resident for the whole session, so the solo versus concurrent memory difference comes only from activation memory and run-time allocations, not from model weights."),
    ("completed_type_in_window / throughput_rps", "Requests of THIS type completed in the window (solo n, co-run n_window) and that count / window_s (requests per second). set_throughput_rps = completed_in_window / window_s (all tenants). delta_throughput_* against the solo throughput of the type at the same level (median over repeats)."),
    ("gpu_util_measured", "n = 1 only: u of solo.csv = mean over the clock samples of the solo window of gpu_load_permille / 1000 (0 to 1). Measured."),
    ("gpu_demand_offered", "n = 2, 3 only: U_A of corun.csv = SUM over the members of A of the solo u at the same level (mean over solo repeats), as computed in closed_loop.py. It is offered demand, not a measured co-run load, can exceed 1, and is not compared with solo: there is no delta column for it."),
    ("Note", "Solo rows have delta = 0 against the median only up to the difference between the pooled solo median and the median of the repeat medians."),
]


def write_cost(out, solo, corun, figs_dir, p_idle=None, mem=None, screen=None):
    d = Path(out) / "1_cost_value"
    rows = build_cost(solo, corun, p_idle, mem, screen)
    for fam in FAMS:
        sheets = [("summary", SUM_COLS, cost_summary(rows, fam))]
        sheets += [(t.replace(":", "_"), COST_COLS, rows.get(t, [])) for t in TYPES_OF[fam]]
        write_workbook(d / f"cost_{fam}.xlsx", sheets, COST_DEFS)
    copied = _copy_figs(figs_dir, d, COST_FIGS)
    fig_cost_delta(rows, d)
    return rows, copied


TITLES = True          # --no-titles switches the plain-English titles off (paper version)
NAMES = {"bert": "BERT", "llama": "LLaMA", "yolo": "YOLO"}


def plain(t):
    """'bert:5' -> 'BERT exit 5'."""
    fam, ex = t.split(":")
    return f"{NAMES[fam]} exit {ex}"


def _big_font():
    F._style()
    plt.rcParams.update({"font.size": 10, "axes.labelsize": 10, "xtick.labelsize": 10, "ytick.labelsize": 10,
                         "legend.fontsize": 10, "axes.titlesize": 11})


def _family_values(rows, fam, key, n, f):
    """(median over the family's exits and sets, {type: median over the sets of that exit})."""
    per = {}
    for t in TYPES_OF[fam]:
        v = [r[key] for r in rows.get(t, []) if r["repeat"] == "median" and r["n_concurrent"] == n
             and r["f_mhz"] == f and r[key] is not None]
        if v:
            per[t] = st.median(v)
    allv = [r[key] for t in TYPES_OF[fam] for r in rows.get(t, []) if r["repeat"] == "median"
            and r["n_concurrent"] == n and r["f_mhz"] == f and r[key] is not None]
    return (st.median(allv) if allv else None), per


def _bar_figure(rows, out, name, title, ylabel, keys, key_labels=None):
    """Grouped bars: x = GPU frequency, bars = n = 2 and n = 3 (and, per n, one bar per key: the second key hatched),
    one panel per family. Bar = median over the exits and sets of the family, dots = the exits."""
    _big_font()
    fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.9), sharey=True)
    nb = len(keys) * 2
    width = 0.8 / nb
    vals = []
    for ax, fam in zip(axes, FAMS):
        col = F.FAMILY_COLOR[fam]
        for gi, f in enumerate(F.LEVELS):
            for ni, n in enumerate((2, 3)):
                for ki, key in enumerate(keys):
                    pos = gi - 0.4 + width * (ni * len(keys) + ki + 0.5)
                    v, per = _family_values(rows, fam, key, n, f)
                    if v is None:
                        continue
                    ax.bar(pos, v, width * 0.92, color=col, alpha=0.5 if n == 2 else 1.0,
                           hatch="//" if ki == 1 else None, edgecolor="black", linewidth=0.5, zorder=2)
                    ax.scatter([pos] * len(per), list(per.values()), s=14, color="white", edgecolor="black",
                               linewidth=0.8, zorder=4)
                    tip = max([v] + list(per.values())) if v >= 0 else min([v] + list(per.values()))
                    ax.annotate(f"{v:+.0f}", (pos, tip), xytext=(0, 4 if v >= 0 else -4), textcoords="offset points",
                                ha="center", va="bottom" if v >= 0 else "top", fontsize=10, zorder=5)
                    vals += [v] + list(per.values())
        ax.axhline(0, color="black", linewidth=1.0, zorder=3)
        ax.set_xticks(range(3), [f"{f} MHz" for f in F.LEVELS])
        ax.set_xlabel("GPU frequency")
        ax.set_title(NAMES[fam], color=col, fontweight="bold")
        ax.grid(axis="y", linewidth=0.3, zorder=0)
    lo, hi = min(vals + [0]), max(vals + [0])
    pad = 0.14 * (hi - lo or 1)
    axes[0].set_ylim(lo - pad, hi + pad)
    axes[0].set_ylabel(ylabel)
    from matplotlib.patches import Patch
    handles = [Patch(facecolor="#777777", alpha=0.5, edgecolor="black", label="2 requests run together"),
               Patch(facecolor="#777777", edgecolor="black", label="3 requests run together"),
               Line2D([], [], marker="o", linestyle="", markerfacecolor="white", markeredgecolor="black",
                      label="each exit of the family"),
               Line2D([], [], color="black", linewidth=1.5, label="zero line: same as running alone")]
    if key_labels:
        handles.insert(0, Patch(facecolor="white", edgecolor="black", label=key_labels[0]))
        handles.insert(1, Patch(facecolor="white", edgecolor="black", hatch="//", label=key_labels[1]))
    fig.legend(handles=handles, loc="lower center", ncol=3, frameon=False,
               bbox_to_anchor=(0.5, -0.04))
    if TITLES:
        fig.suptitle(title, fontsize=13, fontweight="bold")
    fig.tight_layout(rect=(0, 0.13, 1, 0.94 if TITLES else 1))
    return F._save(fig, out, name)


def fig_cost_delta(rows, out):
    made = []
    made += _bar_figure(rows, out, "fig_delta_latency",
                        "How much slower is a request when 2 or 3 run together?",
                        "Change in time per request vs running alone (%)", ("delta_tau_pct",))
    made += _bar_figure(rows, out, "fig_delta_power",
                        "How much more power does the board draw when 2 or 3 run together?",
                        "Change in board power vs running alone (%)", ("delta_p_pct",))
    made += _bar_figure(rows, out, "fig_delta_energy",
                        "How does the energy per request change when 2 or 3 run together?",
                        "Change in energy per request vs running alone (%)",
                        ("delta_energy_dyn_pct", "delta_energy_total_pct"),
                        ("dynamic energy (idle power removed)", "total energy (idle power included)"))
    made += _bar_figure(rows, out, "fig_delta_throughput",
                        "How many fewer requests per second does a request type complete when 2 or 3 run together?",
                        "Change in requests per second vs running alone (%)", ("delta_throughput_pct",))
    return made


# ------------------------------------------------------------------------------------------------ 3 slowdown factor
SLOW_COLS = ["type_id", "n_concurrent", "f_mhz", "co_run_set_A", "cell_id", "split", "repeat", "source",
             "slowdown_s_p50", "tau_solo_p50_s", "tau_corun_p50_s", "delta_tau_s", "delta_tau_pct", "capacity_c"]


def _slow_row(base):
    r = dict(base)
    r["tau_corun_p50_s"] = r["slowdown_s_p50"] * r["tau_solo_p50_s"]
    r["delta_tau_s"] = r["tau_corun_p50_s"] - r["tau_solo_p50_s"]
    r["delta_tau_pct"] = (r["slowdown_s_p50"] - 1.0) * 100.0
    return r


def build_slowdown(corun):
    cap = F.capacity_table(corun)
    out = collections.defaultdict(list)
    for (cid, f, rep), mem in sorted(corun_groups(corun).items()):
        c = sum(1.0 / _f(x["s_p50"]) for x in mem)
        for x in mem:
            out[x["type_id"]].append(_slow_row(dict(
                type_id=x["type_id"], n_concurrent=_n_of(x["A"]), f_mhz=f, co_run_set_A=x["A"], cell_id=cid,
                split=x["split"], repeat=int(rep), source=_src(cid), slowdown_s_p50=_f(x["s_p50"]),
                tau_solo_p50_s=_f(x["tau_solo_p50"]), capacity_c=c)))
    for t in list(out):
        med = []
        grp = collections.defaultdict(lambda: collections.defaultdict(list))
        for r in out[t]:
            grp[(r["n_concurrent"], r["co_run_set_A"], r["f_mhz"])][r["source"]].append(r)
        for (n, A, f), bysrc in sorted(grp.items()):
            rs = _prefer_d(bysrc)
            med.append(_slow_row(dict(
                type_id=t, n_concurrent=n, f_mhz=f, co_run_set_A=A, cell_id="median", split=rs[0]["split"],
                repeat="median", source="d" if bysrc.get("d") else "c",
                slowdown_s_p50=_med([x["slowdown_s_p50"] for x in rs]),
                tau_solo_p50_s=rs[0]["tau_solo_p50_s"], capacity_c=cap[(A, f)])))
        out[t].extend(med)
    return out, cap


SLOW_SUM_COLS = ["family", "n_concurrent", "f_mhz", "median_rows", "slowdown_s_p50", "delta_tau_pct", "capacity_c_min",
                 "capacity_c_max"]


def slow_summary(rows, fam):
    res = []
    for n in (2, 3):
        for f in F.LEVELS:
            rs = [r for t in TYPES_OF[fam] for r in rows.get(t, []) if r["repeat"] == "median"
                  and r["n_concurrent"] == n and r["f_mhz"] == f]
            if rs:
                res.append(dict(family=fam, n_concurrent=n, f_mhz=f, median_rows=len(rs),
                                slowdown_s_p50=_med([r["slowdown_s_p50"] for r in rs]),
                                delta_tau_pct=_med([r["delta_tau_pct"] for r in rs]),
                                capacity_c_min=min(r["capacity_c"] for r in rs),
                                capacity_c_max=max(r["capacity_c"] for r in rs)))
    return res


SLOW_DEFS = [
    ("slowdown_s_p50", "Co-run median service time / solo median service time of the same type at the same frequency (corun.csv s_p50)."),
    ("tau_solo_p50_s", "Pooled solo median service time (s), tau_solo_p50 of corun.csv."),
    ("tau_corun_p50_s", "slowdown_s_p50 x tau_solo_p50_s (s)."),
    ("delta_tau_s / delta_tau_pct", "tau_corun_p50_s - tau_solo_p50_s (s) and (slowdown_s_p50 - 1) x 100."),
    ("capacity_c", "c(A,f) = sum over members of A of 1/s_p50. Row of one repeat: that repeat. Median row: median over repeats, Run 1d rows where Run 1d re-measured the cell (as Table capacity of the paper). Same value on every member row of the set."),
    ("n_concurrent", "Number of members of the co-run set A (2 or 3). n = 1 is solo: slowdown 1, delta 0 (not listed)."),
    ("repeat = median", "Median over repeats of each (type, set, frequency); source d marks Run 1d rows."),
    ("split", "train or heldout cell of the slowdown fit."),
]


def write_slowdown(out, corun, figs_dir):
    d = Path(out) / "3_slowdown_factor"
    rows, cap = build_slowdown(corun)
    for fam in FAMS:
        sheets = [("summary", SLOW_SUM_COLS, slow_summary(rows, fam))]
        sheets += [(t.replace(":", "_"), SLOW_COLS, rows.get(t, [])) for t in TYPES_OF[fam]]
        write_workbook(d / f"slowdown_{fam}.xlsx", sheets, SLOW_DEFS)
    copied = _copy_figs(figs_dir, d, SLOW_FIGS)
    fig_slowdown_by_n(rows, d)
    return rows, cap, copied


EXIT_COLORS = {fam: (F.TYPE_STYLE[a][0], F.TYPE_STYLE[b][0])   # same exit colours as figures.py
               for fam, (a, b) in {"bert": ("bert:5", "bert:23"), "llama": ("llama:3", "llama:15"),
                                   "yolo": ("yolo:1", "yolo:5")}.items()}


def fig_slowdown_by_n(rows, out):
    """One panel per family: slowdown of each request type (exit) at 612, 816 and 1020 MHz, n = 2 and n = 3."""
    _big_font()
    fig, axes = plt.subplots(1, 3, figsize=(13.5, 5.0), sharey=True)
    for ax, fam in zip(axes, FAMS):
        for ti, t in enumerate(TYPES_OF[fam]):
            for n, mk, dx in ((2, "o", -0.09), (3, "^", 0.09)):
                for gi, f in enumerate(F.LEVELS):
                    v = [r["slowdown_s_p50"] for r in rows.get(t, []) if r["repeat"] == "median"
                         and r["n_concurrent"] == n and r["f_mhz"] == f]
                    if v:
                        x = gi + (ti - 0.5) * 0.28 + dx * 0.9
                        ax.scatter([x + (k - (len(v) - 1) / 2) * 0.012 for k in range(len(v))], v, marker=mk, s=42,
                                   color=EXIT_COLORS[fam][ti], edgecolor="black", linewidth=0.5, alpha=0.9,
                                   zorder=3)
        ax.axhline(1, color="black", linewidth=1.0, zorder=2)
        ax.annotate("running alone = 1", (0.01, 1), xycoords=("axes fraction", "data"), xytext=(0, 3),
                    textcoords="offset points", fontsize=10, va="bottom",
                    bbox=dict(facecolor="white", edgecolor="none", alpha=0.8, pad=1))
        ax.set_xticks(range(3), [f"{f} MHz" for f in F.LEVELS])
        ax.set_xlim(-0.5, 2.5)
        ax.set_xlabel("GPU frequency")
        ax.set_title(NAMES[fam], color=F.FAMILY_COLOR[fam], fontweight="bold")
        ax.grid(axis="y", linewidth=0.3, zorder=0)
        ax.legend(handles=[Line2D([], [], marker="s", linestyle="", color=EXIT_COLORS[fam][i], markeredgecolor="black",
                                  label=plain(t)) for i, t in enumerate(TYPES_OF[fam])], loc="upper right",
                  frameon=False)
    axes[0].set_ylabel("Slowdown (x, running alone = 1)")
    fig.legend(handles=[Line2D([], [], marker="o", linestyle="", markerfacecolor="#999999", markeredgecolor="black",
                               label="2 requests run together (one dot per set)"),
                        Line2D([], [], marker="^", linestyle="", markerfacecolor="#999999", markeredgecolor="black",
                               label="3 requests run together (one dot per set)")],
               loc="lower center", ncol=2, frameon=False, bbox_to_anchor=(0.5, -0.02))
    if TITLES:
        fig.suptitle("How many times slower does a request run when it shares the GPU?", fontsize=13,
                     fontweight="bold")
    fig.tight_layout(rect=(0, 0.07, 1, 0.94 if TITLES else 1))
    return F._save(fig, out, "fig_slowdown_by_n")


# ------------------------------------------------------------------------------------------------ 2 warm-up
def load_sessions(dirs):
    """-> {(family, state): dict(ready_s, first_max_s, first_by_exit{exit: s}, load_s, energy_j)}."""
    res = {}
    for d in dirs or []:
        d = Path(d)
        sj = d / "session.json"
        if not sj.is_file():
            continue
        state = "cold" if d.name.endswith("_cold") else "warm" if d.name.endswith("_warm") else None
        sess = F.json.load(open(sj))
        for fam, srv in (sess.get("servers") or {}).items():
            if state is None or "t_spawn_ns" not in srv or "t_ready_ns" not in srv:
                continue
            t0, t1 = srv["t_spawn_ns"], srv["t_ready_ns"]
            wu = srv.get("ready", {}).get("warmup", {})
            fi = {str(k): float(v) for k, v in (wu.get("first_infer_s") or {}).items()}
            res[(fam, state)] = dict(ready_s=(t1 - t0) / 1e9, first_max_s=max(fi.values()) if fi else None,
                                     first_by_exit=fi, load_s=wu.get("load_s"),
                                     energy_j=F._energy_j(d / "telemetry" / "power.csv", t0, t1), session=d.name)
    return res


START_COLS = ["family", "metric", "cold", "warm", "delta_warm_minus_cold", "delta_pct", "ratio_cold_over_warm"]
WARM_ROW_COLS = ["session_id", "family", "t_load_s", "first_infer_s_max", "warm_energy_j", "cache"]
TYPE_COLS = ["type_id", "metric", "cold", "warm", "delta_warm_minus_cold", "delta_pct", "ratio_cold_over_warm"]


def _cw(key, cold, warm, **kw):
    ratio = cold / warm if cold is not None and warm else None
    return dict(cold=cold, warm=warm, delta_warm_minus_cold=_diff(warm, cold), delta_pct=_pct(warm, cold),
                ratio_cold_over_warm=ratio, **kw)


def startup_rows(res, fam):
    out = []
    for key, name in (("ready_s", "spawn_to_ready_s"), ("first_max_s", "first_inference_s (longest warmed exit)"),
                      ("load_s", "model_load_s (inside start-up)"),
                      ("energy_j", "startup_energy_j (integral of VDD_IN, spawn to ready)")):
        c, w = res.get((fam, "cold"), {}).get(key), res.get((fam, "warm"), {}).get(key)
        if c is not None or w is not None:
            out.append(_cw(key, c, w, family=fam, metric=name))
    return out


def type_rows(res, fam):
    out = []
    c, w = res.get((fam, "cold"), {}).get("first_by_exit", {}), res.get((fam, "warm"), {}).get("first_by_exit", {})
    for t in TYPES_OF[fam]:
        ex = t.split(":")[1]
        if ex in c or ex in w:
            out.append(_cw("first", c.get(ex), w.get(ex), type_id=t, metric="first_inference_s"))
    return out


WARM_DEFS = [
    ("cold / warm", "Run 1c: one session per family and cache state, the family started alone at 1020 MHz. Cold = empty compilation caches, warm = caches filled by the cold start."),
    ("spawn_to_ready_s", "(t_ready_ns - t_spawn_ns) / 1e9 of session.json (s). Table tab:startup of the paper."),
    ("first_inference_s", "First inference after load, per warmed exit (session.json warmup.first_infer_s). The family row uses the longest exit, as in tab:startup."),
    ("startup_energy_j", "Integral (trapezoid) of board power VDD_IN (vdd_in_mw of telemetry/power.csv) from spawn to ready (J)."),
    ("delta_warm_minus_cold", "warm - cold, same unit as the metric. Negative = warm start is cheaper."),
    ("delta_pct", "(warm / cold - 1) x 100."),
    ("sheet warmup_csv", "DIFFERENT DEFINITION: rows of warmup.csv from other sessions. t_load_s is the model-load phase only, warm_energy_j is the energy of the load phase. Not comparable with spawn-to-ready."),
    ("per-type sheets", "A start-up warms every exit of the family. A type has no start-up of its own: only its first inference is separate."),
]


def write_warmup(out, res, warm_rows, figs_dir):
    d = Path(out) / "2_warmup"
    for fam in FAMS:
        sheets = [("summary", START_COLS, startup_rows(res, fam))]
        trs = type_rows(res, fam)
        sheets += [(t.replace(":", "_"), TYPE_COLS, [r for r in trs if r["type_id"] == t]) for t in TYPES_OF[fam]]
        wr = [dict(session_id=r["session_id"], family=r["family"], t_load_s=_f(r["t_load_s"]),
                   first_infer_s_max=_f(r["first_infer_s_max"]), warm_energy_j=_f(r["warm_energy_j"]),
                   cache=r["cache"]) for r in warm_rows if r["family"] == fam]
        sheets.append(("warmup_csv_other_definition", WARM_ROW_COLS, wr))
        write_workbook(d / f"warmup_{fam}.xlsx", sheets, WARM_DEFS)
    copied = _copy_figs(figs_dir, d, WARM_FIGS)
    return copied


# ------------------------------------------------------------------------------------------------ figures copy
def _copy_figs(src, dst, names):
    done = []
    Path(dst).mkdir(parents=True, exist_ok=True)
    for n in names:
        for ext in (".png", ".pdf"):
            p = Path(src) / (n + ext) if src else None
            if p and p.is_file():
                shutil.copy2(p, Path(dst) / p.name)
                done.append(p.name)
            else:
                print(f"[update_pack] figure {n}{ext} not found in {src}")
    return done


# ------------------------------------------------------------------------------------------------ checks + README
def _count_rows(path):
    wb = load_workbook(path)
    return {ws.title: ws.max_row - 1 for ws in wb.worksheets}


def run_checks(out, solo, corun, cap, res, cost_rows=None, gam=None, p_idle=None):
    lines = []

    def add(ok, text):
        lines.append(f"| {'PASS' if ok else 'FAIL'} | {text} |")

    # capacity ranges per group and level
    grp = collections.defaultdict(list)
    for (A, f), c in cap.items():
        grp[(F.set_group(A), f)].append(c)
    for g, byf in PAPER_CAP_GROUPS.items():
        for f, (lo, hi) in byf.items():
            v = grp.get((g, f))
            if not v:
                lines.append(f"| n.a. | capacity {g}, {f} MHz: no cells |")
                continue
            add(round(min(v), 2) == lo and round(max(v), 2) == hi,
                f"capacity {g}, {f} MHz: data {min(v):.3f} to {max(v):.3f} (rounded {min(v):.2f} to {max(v):.2f}), "
                f"paper {lo:.2f} to {hi:.2f}, {len(v)} cells")
    ed = [c for (A, f), c in cap.items() if F.set_group(A) == "encoder and detector"]
    allc = list(cap.values())
    if ed:
        add(round(min(ed), 2) == 0.90 and round(max(ed), 2) == 1.03, f"capacity without decoder {min(ed):.3f} to {max(ed):.3f}, paper 0.90 to 1.03")
    if allc:
        add(round(max(allc), 2) == 1.34, f"capacity maximum with decoder {max(allc):.3f}, paper 1.34")
    # start-up table
    for fam, (cs, ws, cf, wf) in PAPER_STARTUP.items():
        c, w = res.get((fam, "cold")), res.get((fam, "warm"))
        if not c or not w:
            lines.append(f"| n.a. | start-up {fam}: no Run 1c session |")
            continue
        got = (round(c["ready_s"], 1), round(w["ready_s"], 1), round(c["first_max_s"], 1), round(w["first_max_s"], 1))
        add(got == (cs, ws, cf, wf), f"start-up {fam} cold/warm ready and first inference (s): data {got}, paper {(cs, ws, cf, wf)}")
    # solo medians: workbook value against the direct median of solo.csv
    direct = collections.defaultdict(list)
    for r in solo:
        direct[(r["type_id"], int(_f(r["f_mhz"])))].append(_f(r["tau_p50"]))
    bad, tot = 0, 0
    for fam in FAMS:
        p = Path(out) / "1_cost_value" / f"cost_{fam}.xlsx"
        if not p.is_file():
            continue
        wb = load_workbook(p)
        for t in TYPES_OF[fam]:
            if t.replace(":", "_") not in wb.sheetnames:
                continue
            ws = wb[t.replace(":", "_")]
            hdr = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                r = dict(zip(hdr, row))
                if r["n_concurrent"] == 1 and r["repeat"] == "median":
                    tot += 1
                    if abs(r["tau_p50_s"] - st.median(direct[(t, r["f_mhz"])])) > 1e-9:
                        bad += 1
    add(tot > 0 and bad == 0, f"solo median tau_p50 in the workbooks equals the median of solo.csv in {tot - bad} of {tot} (type, frequency) cells")
    # dynamic energy delta against gamma: gamma < 1 means more dynamic energy per unit of work than serial
    if cost_rows is not None and gam:
        sl = set_level_dyn_delta(solo, corun, p_idle)
        ks = [k for k in sl if k in gam]
        a = [sl[k] for k in ks]
        b = [(1.0 / gam[k] - 1.0) * 100.0 for k in ks]
        add(bool(ks) and all(x > 0 for x in a) and len(ks) == sum(g < 1 for g in gam.values()), f"OFFICIAL sign check, set level: dynamic delta (window dynamic energy / sum of n_window x solo dynamic energy per request, minus 1) > 0 in {sum(x > 0 for x in a)} of {len(ks)} cases, range {min(a):.1f} % to {max(a):.1f} %; (1/gamma - 1) x 100 ranges {min(b):.1f} % to {max(b):.1f} %; difference median {st.median([x - y for x, y in zip(a, b)]):.1f} points, largest {max(abs(x - y) for x, y in zip(a, b)):.1f} points; correlation {float(np.corrcoef(a, b)[0, 1]):.2f}")
        for fam in FAMS:
            for key, lab in (("delta_energy_dyn_pct", "dynamic"), ("delta_energy_total_pct", "total incl. idle")):
                v = [r[key] for t in TYPES_OF[fam] for r in cost_rows.get(t, [])
                     if r["repeat"] == "median" and r["n_concurrent"] > 1 and r[key] is not None]
                if v:
                    lines.append(f"| info | {fam}: {key} ({lab}) over all median rows (n = 2, 3) from {min(v):.1f} % to {max(v):.1f} %, median {st.median(v):.1f} % |")
    else:
        lines.append("| n.a. | sign check against gamma: no idle.csv |")
    # screen.csv join and the new column groups
    n_all = STATS["matched"] + STATS["unmatched"]
    add(n_all > 0 and STATS["unmatched"] == 0,
        f"screen.csv join (key cell or type, level, repeat rank in session order): {STATS['matched']} of {n_all} solo and co-run rows matched, {STATS['unmatched']} not matched")
    if cost_rows is not None:
        for fam in FAMS:
            for key, lab, unit in (("delta_mem_available_mb", "min MemAvailable vs solo", "MB"),
                                   ("delta_throughput_pct", "throughput vs solo", "%"),
                                   ("gpu_util_measured", "measured u, solo", ""),
                                   ("gpu_demand_offered", "offered demand U_A", "")):
                for n in (1, 2, 3):
                    v = [r[key] for t in TYPES_OF[fam] for r in cost_rows.get(t, [])
                         if r["repeat"] == "median" and r["n_concurrent"] == n and r[key] is not None]
                    if v and not (key in ("delta_mem_available_mb", "delta_throughput_pct") and n == 1):
                        lines.append(f"| info | {fam}, n = {n}: {key} ({lab}) from {min(v):.3g} {unit} to {max(v):.3g} {unit}, median {st.median(v):.3g} {unit} |")
        tot = bad = 0
        for t, rs in cost_rows.items():
            for r in rs:
                if r["repeat"] != "median":
                    tot += 1
                    bad += abs(r["throughput_rps"] * r["window_s"] - r["completed_type_in_window"]) > 1e-6
        add(tot > 0 and bad == 0, f"throughput_rps x window_s equals the completed requests of the type in {tot - bad} of {tot} rows")
    # reload every workbook
    counts = {}
    for p in sorted(Path(out).glob("*/*.xlsx")):
        try:
            counts[f"{p.parent.name}/{p.name}"] = _count_rows(p)
            add(True, f"{p.parent.name}/{p.name} opens; rows per sheet {counts[f'{p.parent.name}/{p.name}']}")
        except Exception as e:  # noqa: BLE001
            add(False, f"{p.parent.name}/{p.name} does not open: {e}")
    return lines, counts


README = """# update_17: cost value, warm-up and slowdown factor ({date})

Owner (verbatim): "Buat foleer terpisah. Yg pertama itu ngukur cost value semuanya kedua ngukur warmup dan yg ke tiga ngukur slowdown factor. Ini buat jasi uodate_17 isinya ada figure serta ada excel, utj xlsx minta tlg dibuatkan delta dengan nilai delta swbagai perbandingan antara solo run serta n conccurwnt run. Buat xlsx per model."

## Reading of the request (owner: correct this if wrong)

| Term | Reading |
|---|---|
| per model | per model family: bert, llama, yolo. One workbook per family and folder. One sheet per request type of the family (bert:5, bert:23, llama:3, llama:15, yolo:1, yolo:5), plus "summary" and "definitions" |
| n concurrent run | the request runs while n requests are in service in total. n = 1 solo, n = 2 pair, n = 3 triple (co-run set A of corun.csv, n = number of members) |
| delta | concurrent value minus solo value at the same GPU frequency. delta % = (concurrent / solo - 1) x 100 |
| warm-up delta | no concurrent run exists: delta = warm minus cold |

## Contents

```
update_17/
  README.md
  1_cost_value/        cost_{{bert,llama,yolo}}.xlsx   fig_solo_time fig_power fig_gamma fig_run2_energy_attain fig_delta_latency fig_delta_power fig_delta_energy fig_delta_throughput
  2_warmup/            warmup_{{bert,llama,yolo}}.xlsx fig_warmup fig3d_warmup
  3_slowdown_factor/   slowdown_{{bert,llama,yolo}}.xlsx fig_capacity fig3d_capacity fig_slowdown_fit fig_slowdown_by_n
```

Each figure exists as PNG and PDF. fig_delta_latency, fig_delta_power, fig_delta_energy, fig_delta_throughput (one question each, bars per GPU frequency, n = 2 and n = 3, one panel per family, dots = the exits; energy has dynamic solid and total hatched) and fig_slowdown_by_n are new (drawn by update_pack.py, option --no-titles removes the titles). Slowdown, delta and energy numbers are unchanged. The others are copies of the paper figures.

## Data sources

| Source | Path |
|---|---|
| solo.csv, corun.csv, warmup.csv | {run1_csv} |
| Run 1c sessions (cold, warm) | {run1c} |
| copied figures | {figures} |
| code | exp2/analyze/update_pack.py (reuses exp2/analyze/figures.py); command: python -m exp2.analyze.update_pack --run1-csv ... --run1c ... --figures ... --out ... |

## Columns of the cost workbooks (1_cost_value)

| Group | Columns |
|---|---|
| Time, power | tau_p50_s, tau_p95_s, p_board_w, window_s, delta_tau_*, delta_p_* |
| Energy | energy_window_j, p_idle_w, dyn_ratio_D_over_S, energy_total_per_request_j, energy_dyn_per_request_j, delta_energy_total_*, delta_energy_dyn_* |
| Memory | resident_mb, act_mb, act_sum_mb, delta_act_mb, delta_act_pct, min_mem_available_mb, delta_mem_available_mb, delta_mem_available_pct |
| Throughput | completed_type_in_window, throughput_rps, set_throughput_rps, delta_throughput_rps, delta_throughput_pct |
| GPU utilization | gpu_util_measured (u of solo.csv, n = 1), gpu_demand_offered (U_A of corun.csv, n = 2, 3), no delta |

Memory: in Run 1 all three families stay resident for the whole session, so the solo versus concurrent memory difference comes only from activation memory and run-time allocations, not from model weights. The memory columns come from mem.csv and screen.csv. A screen.csv row is matched to a solo or co-run row by condition id (formats solo_bert-5_f612_r0, c001_f612_r2, d_c002_f816_r0, d_solo_llama-15_f612_r0): key = cell or type and level, and the k-th smallest repeat takes the k-th kept row in session_id order. The repeat order within a key follows session_id order (an assumption); the workbook medians over repeats do not depend on it. screen.csv has no swap or page-fault values for Run 1. The check table gives the match count.

GPU utilization caution: U_A of corun.csv (column gpu_demand_offered) is the sum over the set members of the solo u at the same level (closed_loop.py), so it is offered demand and not a measured co-run load. It can exceed 1 and has no delta against solo.

## Definitions that need care

| Item | Definition |
|---|---|
| Co-run service time | s_p50 x tau_solo_p50 (corun.csv). tau_solo_p50 is the pooled solo median of corun.csv. It differs slightly from the median of the per-repeat medians of solo.csv |
| Total energy per request (co-run) | board energy including idle, busy period. W_window = sum over tenants of n_window x tau_solo_p50 (solo-equivalent work, s). e_work = energy_j / W_window. Energy per request of type j = e_work x tau_solo_p50 of j. Solo: p_board_w x window_s / n (median over repeats). Columns energy_total_per_request_j, delta_energy_total_j, delta_energy_total_pct |
| Dynamic energy per request | idle subtracted with the warm idle power P_idle(f) at the same level, models resident (idle.csv kind resident, the same values as p_idle_w of power.json and of gamma). Co-run: d_solo_j x D_window / S_window with D_window = energy_j - P_idle x window_s and S_window = sum over tenants of n_window x d_solo (attribution in proportion to solo dynamic energy; the delta is the same for every member of the window). Solo: (p_board_w - P_idle) x window_s / n. Columns energy_dyn_per_request_j, delta_energy_dyn_j, delta_energy_dyn_pct |
| Total delta below zero at 1020 MHz | the capacity of a set is above 1 there, so the set completes more work per second and shares the idle power over more work, and the total energy per request can fall below solo. The dynamic energy has no idle part and rises (gamma below 1 means co-running costs more dynamic energy per unit of work than serial execution), so the dynamic delta stays positive |
| Removed measure | energy_j / requests of all tenants mixes cheap and expensive types and is not comparable to solo. It is not in the workbooks |
| Power of a co-run row | board power of the whole set, not of one tenant |
| Capacity c(A,f) | sum over members of 1/s_p50. Median rows follow the paper (Run 1d rows replace Run 1c rows where Run 1d re-measured the cell) |
| Start-up energy | integral of VDD_IN from spawn to ready (Run 1c telemetry). The warmup.csv sheet uses another definition (load phase only, other sessions) and is labelled so |
| Frequencies | 612, 816, 1020 MHz for co-run. Solo also has 408 and 714 MHz (no delta there, no co-run at that level) |

## Checks

| Result | Check |
|---|---|
{checks}

Reading of the checks: PASS means the number reproduces the paper value after rounding. FAIL lists data and paper value side by side.

The energy of one request type inside a shared GPU is not separately measurable, so every per-type energy value is an attribution. Total uses tau-proportional attribution, dynamic uses solo-dynamic-energy attribution, and dynamic is the one comparable to gamma in the paper.

On the sign check: gamma(A,f) = (sum_i dP_i g_i) / (P(A,f) - P_idle(f)) with dP_i = P(i alone,f) - P_idle(f) (method.tex, Eq. eq:gamma). Gamma below 1 means co-running costs more dynamic energy per unit of work than serial execution, so the dynamic delta is expected above zero, and near (1/gamma - 1) x 100 at set level.
"""


def write_readme(out, args, lines):
    txt = README.format(date=DATE, run1_csv=args.run1_csv, run1c="; ".join(map(str, args.run1c)),
                        figures=args.figures, checks="\n".join(lines))
    (Path(out) / "README.md").write_text(txt, encoding="utf-8")


# ------------------------------------------------------------------------------------------------ main
def build(run1_csv, run1c, figures, out, args=None, titles=True):
    global TITLES
    TITLES = titles
    solo, corun, warm_rows, p_idle, mem, screen = load_inputs(run1_csv)
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    cost_rows, _ = write_cost(out, solo, corun, figures, p_idle, mem, screen)
    _, cap, _ = write_slowdown(out, corun, figures)
    res = load_sessions(run1c)
    write_warmup(out, res, warm_rows, figures)
    gam = None
    if p_idle:
        gam = F.gamma_table(corun, solo, p_idle)
    lines, counts = run_checks(out, solo, corun, cap, res, cost_rows, gam, p_idle)
    if args is None:
        args = argparse.Namespace(run1_csv=run1_csv, run1c=run1c, figures=figures)
    write_readme(out, args, lines)
    return lines, counts


def main(argv=None):
    ap = argparse.ArgumentParser(prog="exp2.analyze.update_pack")
    ap.add_argument("--run1-csv", required=True)
    ap.add_argument("--run1c", nargs="*", default=[])
    ap.add_argument("--figures", default=None)
    ap.add_argument("--out", required=True)
    ap.add_argument("--no-titles", action="store_true", help="draw the new figures without the plain-English title (paper)")
    a = ap.parse_args(argv)
    lines, counts = build(a.run1_csv, a.run1c, a.figures, a.out, a, titles=not a.no_titles)
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
